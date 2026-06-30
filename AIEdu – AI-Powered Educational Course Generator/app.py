import os, json, sqlite3, uuid, io, re
from datetime import datetime
from flask import Flask, render_template, request, redirect, url_for, send_file, flash
from dotenv import load_dotenv
from openai import OpenAI
from pptx import Presentation

load_dotenv()
app = Flask(__name__, instance_relative_config=True, static_folder="static")
app.secret_key = os.getenv("FLASK_SECRET_KEY", "dev_secret_SECRET_KEY")
os.makedirs(app.instance_path, exist_ok=True)
DB_PATH = os.path.join(app.instance_path, "courses.sqlite3")

client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))

# ---------------- DB ----------------
def init_db():
    with sqlite3.connect(DB_PATH) as conn:
        c = conn.cursor()
        c.execute("""
        CREATE TABLE IF NOT EXISTS courses (
            id TEXT PRIMARY KEY,
            created_at TEXT,
            topic TEXT,
            academic_level TEXT,
            language TEXT,
            duration_minutes INTEGER,
            num_chapters INTEGER,
            payload_json TEXT
        )""")
        conn.commit()

def save_course(payload, form_meta):
    course_id = str(uuid.uuid4())
    with sqlite3.connect(DB_PATH) as conn:
        c = conn.cursor()
        c.execute("""INSERT INTO courses
            (id, created_at, topic, academic_level, language, duration_minutes, num_chapters, payload_json)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            (
                course_id,
                datetime.utcnow().isoformat(),
                form_meta["topic"],
                form_meta["academic_level"],
                form_meta["language"],
                int(form_meta["duration_minutes"]),
                int(form_meta["num_chapters"]),
                json.dumps(payload, ensure_ascii=False),
            ),
        )
        conn.commit()
    return course_id

def get_course(course_id):
    with sqlite3.connect(DB_PATH) as conn:
        c = conn.cursor()
        c.execute("""SELECT id, created_at, topic, academic_level, language,
                            duration_minutes, num_chapters, payload_json
                     FROM courses WHERE id=?""", (course_id,))
        row = c.fetchone()
        if not row: return None
        return {
            "id": row[0], "created_at": row[1], "topic": row[2],
            "academic_level": row[3], "language": row[4],
            "duration_minutes": row[5], "num_chapters": row[6],
            "payload": json.loads(row[7])
        }

def list_courses():
    with sqlite3.connect(DB_PATH) as conn:
        c = conn.cursor()
        c.execute("""SELECT id, created_at, topic, academic_level, language,
                            duration_minutes, num_chapters
                     FROM courses ORDER BY datetime(created_at) DESC""")
        rows = c.fetchall()
        return [
            {
                "id": r[0], "created_at": r[1], "topic": r[2],
                "academic_level": r[3], "language": r[4],
                "duration_minutes": r[5], "num_chapters": r[6]
            } for r in rows
        ]

# ---------------- Quiz prompt (professional standard) ----------------
def mcq_prompt_guidelines():
    return (
        "MCQ Generation Standard:\\n"
        "- Source of truth: Only the lesson's content you just wrote. Do not invent facts.\\n"
        "- Quantity: 4–8 multiple-choice questions per lesson.\\n"
        "- Cognitive range: Mix of recall, understanding, application, and analysis (Bloom's levels).\\n"
        "- Structure: Each MCQ has: 'question' (clear and concise), 'options' (3–6 choices), and a single correct answer.\\n"
        "- Options: Plausible, mutually exclusive, similar length/structure; avoid clues like 'always/never'; avoid 'All of the above'/'None of the above'.\\n"
        "- Wording: Avoid ambiguous negatives; avoid trick questions; reflect the terminology used in the lesson.\\n"
        "- Correctness: Exactly one correct option; set 'answer_index' to the zero-based index of the correct option.\\n"
        "- Coverage: Collectively span key concepts, definitions, examples, and common misconceptions addressed in the lesson.\\n"
        "- Language & level: Match the requested language and academic level.\\n"
    )

# ---------------- JSON helpers (fallback when response_format is unavailable) ----------------
def extract_json_block(text: str):
    """Try to robustly extract the first valid top-level JSON object from text."""
    fence = re.search(r"```(?:json)?\\s*(\\{[\\s\\S]*?\\})\\s*```", text, re.IGNORECASE)
    candidate = fence.group(1) if fence else None
    if not candidate:
        start, end = text.find('{'), text.rfind('}')
        if start != -1 and end != -1 and end > start:
            candidate = text[start:end+1]
    if not candidate:
        raise ValueError("No JSON object found in model output.")
    candidate = re.sub(r",\\s*([}\\]])", r"\\1", candidate)
    return json.loads(candidate)

# ---------------- OpenAI generation ----------------
def build_course_schema_description():
    return (
        "Return a SINGLE JSON object with this shape:\\n"
        "{\\n"
        "  \"course_title\": \"...\",\\n"
        "  \"course_description\": \"...\",\\n"
        "  \"academic_level\": \"...\",\\n"
        "  \"language\": \"...\",\\n"
        "  \"duration_minutes\": 90,\\n"
        "  \"outline\": [ { \"chapter_title\": \"...\", \"lessons\": [\"...\", \"...\"] }, ... ],\\n"
        "  \"chapters\": [\\n"
        "     { \"title\": \"...\", \"description\": \"...\", \"lessons\": [\\n"
        "        { \"title\": \"...\", \"summary\": \"...\", \"content\": \"...\",\\n"
        "          \"learning_objectives\": [\"...\", \"...\"],\\n"
        "          \"quiz\": [ { \"question\": \"...\", \"options\": [\"...\", \"...\"], \"answer_index\": 0 }, ... ]\\n"
        "        }\\n"
        "     ] }\\n"
        "  ]\\n"
        "}\\n"
        "- Use double quotes for all JSON keys and strings.\\n"
        "- Do not include any commentary outside the JSON.\\n"
    )

def generate_course_with_openai(topic, academic_level, language, duration_minutes, num_chapters):
    schema_text = build_course_schema_description()
    instructions = (
        "You are an expert instructional designer. Create academically rigorous, age-appropriate course "
        "content aligned to the requested level and language. "
        "First, produce an 'outline' (chapter titles with lesson titles only), then full 'chapters' with rich content. "
        "Quizzes must be derived strictly from lesson content.\\n\\n"
        + mcq_prompt_guidelines() + "\\n\\n" + schema_text
    )
    user_prompt = f"""
Topic: {topic}
Academic level: {academic_level}
Language: {language}
Course duration (minutes): {duration_minutes}
Desired number of chapters: {num_chapters}

Rules:
- The 'outline' must list {num_chapters} chapter titles with 2–6 lesson titles each.
- 'chapters' must follow the same order as 'outline' and expand each lesson with summary, content, objectives, and MCQ quiz.
- Content should be multi-paragraph, markdown-friendly prose. Avoid code blocks unless necessary.
- Quizzes: 4–8 MCQs per lesson; provide options and one correct answer_index.
- Use terminology and difficulty suitable for: {academic_level}, language: {language}.
- IMPORTANT: Respond with ONLY the JSON object, no extra text.
"""
    try:
        resp = client.responses.create(
            model=os.getenv("OPENAI_MODEL","gpt-4o-mini"),
            temperature=0.7,
            max_output_tokens=8000,
            instructions=instructions,
            input=user_prompt,
            response_format={"type": "json_object"}
        )
        text = resp.output_text
        return json.loads(text)
    except TypeError:
        resp = client.responses.create(
            model=os.getenv("OPENAI_MODEL","gpt-4o-mini"),
            temperature=0.7,
            max_output_tokens=8000,
            instructions=instructions,
            input=user_prompt
        )
        text = resp.output_text
        return extract_json_block(text)

# ---------------- PPTX ----------------
def course_to_pptx(course):
    prs = Presentation()
    title_layout = prs.slide_layouts[0]
    content_layout = prs.slide_layouts[1]

    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = course["course_title"]
    slide.placeholders[1].text = course.get("course_description","")

    slide = prs.slides.add_slide(content_layout)
    slide.shapes.title.text = "Course Overview"
    tf = slide.placeholders[1].text_frame
    tf.clear()
    for k in ["academic_level", "language", "duration_minutes"]:
        p = tf.add_paragraph(); p.text = f"{k.replace('_',' ').title()}: {course.get(k)}"; p.level = 0

    slide = prs.slides.add_slide(content_layout)
    slide.shapes.title.text = "Outline"
    tf = slide.placeholders[1].text_frame
    tf.clear()
    for ch in course.get("outline", []):
        p = tf.add_paragraph(); p.text = f"• {ch['chapter_title']}"; p.level = 0
        for ls in ch["lessons"]:
            s = tf.add_paragraph(); s.text = f"- {ls}"; s.level = 1

    for ch in course["chapters"]:
        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = ch["title"]
        slide.placeholders[1].text = ch.get("description","")

        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = f"{ch['title']} — Lessons"
        tf = slide.placeholders[1].text_frame; tf.clear()
        for ls in ch["lessons"]:
            p = tf.add_paragraph(); p.text = f"• {ls['title']}: {ls.get('summary','')}"; p.level = 0

        for ls in ch["lessons"]:
            slide = prs.slides.add_slide(content_layout)
            slide.shapes.title.text = f"Lesson: {ls['title']} — Objectives"
            tf = slide.placeholders[1].text_frame; tf.clear()
            for obj in ls.get("learning_objectives", []):
                p = tf.add_paragraph(); p.text = f"• {obj}"; p.level = 0

            slide = prs.slides.add_slide(content_layout)
            slide.shapes.title.text = f"Lesson: {ls['title']} — Content"
            tf = slide.placeholders[1].text_frame; tf.clear()
            for para in (ls.get("content","").split("\n")):
                para = para.strip()
                if not para: continue
                p = tf.add_paragraph(); p.text = para; p.level = 0

            if ls.get("quiz"):
                slide = prs.slides.add_slide(content_layout)
                slide.shapes.title.text = f"Lesson: {ls['title']} — Quiz"
                tf = slide.placeholders[1].text_frame; tf.clear()
                for i, q in enumerate(ls["quiz"][:6], start=1):
                    p = tf.add_paragraph()
                    correct = q["options"][q["answer_index"]]
                    p.text = f"Q{i}. {q['question']}  (Ans: {correct})"
                    p.level = 0

    bio = io.BytesIO(); prs.save(bio); bio.seek(0); return bio

# ---------------- Routes ----------------
@app.route("/")
def home():
    init_db()
    courses = list_courses()
    return render_template("index.html", courses=courses, html_dir="ltr")

@app.route("/new")
def new_course():
    return render_template("new_course.html", html_dir="ltr")

@app.route("/generate", methods=["POST"])
def generate():
    topic = request.form.get("topic","").strip()
    academic_level = request.form.get("academic_level","").strip()
    language = request.form.get("language","English").strip()
    duration_minutes = int(request.form.get("duration_minutes","90"))
    num_chapters = int(request.form.get("num_chapters","6"))

    if duration_minutes < 30 or duration_minutes > 240:
        flash("Duration must be between 30 and 240 minutes.", "error")
        return redirect(url_for("new_course"))

    if not topic:
        flash("Please enter a course topic.", "error")
        return redirect(url_for("new_course"))

    try:
        course_json = generate_course_with_openai(topic, academic_level, language, duration_minutes, num_chapters)
        course_json.setdefault("academic_level", academic_level)
        course_json.setdefault("language", language)
        course_json.setdefault("duration_minutes", duration_minutes)

        course_id = save_course(course_json, {
            "topic": topic, "academic_level": academic_level, "language": language,
            "duration_minutes": duration_minutes, "num_chapters": num_chapters
        })
        return redirect(url_for("view_course", course_id=course_id))
    except Exception as e:
        flash(f"Generation failed: {e}", "error")
        return redirect(url_for("new_course"))

@app.route("/course/<course_id>")
def view_course(course_id):
    course = get_course(course_id)
    if not course:
        flash("Course not found.", "error")
        return redirect(url_for("home"))
    html_dir = "rtl" if (course["payload"].get("language","").lower().startswith("arab")) else "ltr"
    return render_template("course.html", course=course, html_dir=html_dir)

@app.route("/quiz/<course_id>/<int:chapter_idx>", methods=["GET","POST"])
def quiz(course_id, chapter_idx):
    course = get_course(course_id)
    if not course: 
        flash("Course not found.","error"); return redirect(url_for("home"))
    chapters = course["payload"]["chapters"]
    if chapter_idx < 0 or chapter_idx >= len(chapters):
        flash("Invalid chapter.","error"); return redirect(url_for("view_course", course_id=course_id))
    chapter = chapters[chapter_idx]

    items = []
    for li, ls in enumerate(chapter["lessons"]):
        for qi, q in enumerate(ls.get("quiz", [])):
            items.append({
                "lesson_index": li, "q_index": qi,
                "question": q["question"], "options": q["options"],
                "answer_index": q["answer_index"]
            })

    score, selected = None, {}
    if request.method == "POST":
        correct = 0
        for i in range(len(items)):
            key = f"q{i}"; val = request.form.get(key)
            selected[str(i)] = val
            if val is not None and int(val) == items[i]["answer_index"]: correct += 1
        score = {"correct": correct, "total": len(items)}

    html_dir = "rtl" if (course["payload"].get("language","").lower().startswith("arab")) else "ltr"
    return render_template("quiz.html", course=course, chapter=chapter, items=items,
                           chapter_idx=chapter_idx, score=score, selected=selected, html_dir=html_dir)

@app.route("/download/<course_id>.pptx")
def download_pptx(course_id):
    course = get_course(course_id)
    if not course:
        flash("Course not found.","error"); return redirect(url_for("home"))
    bio = course_to_pptx(course["payload"])
    filename = f"{course['payload']['course_title']}.pptx"
    return send_file(bio, as_attachment=True, download_name=filename,
                     mimetype="application/vnd.openxmlformats-officedocument.presentationml.presentation")

if __name__ == "__main__":
    init_db()
    app.run(debug=True)
